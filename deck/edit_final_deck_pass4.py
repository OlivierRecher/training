# -*- coding: utf-8 -*-
"""Fourth pass: slide 2 restyled. The three equal-height cards left two thirds of
their own space empty and stranded the callout inside the third one. Replaced by
an asymmetric split: the question stated once, large, on the left; the baseline
and the reporting rule as two flat bands on the right, in the deck's own
ground + left-accent-bar grammar; the reading moved to a full-width band."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK = "../Présentation Stage - Finale.pptx"

RULE, BASE, BODY, TITLE, SECOND = "387190", "4E97B0", "33363D", "232959", "47515C"
NEUTRAL, BASEGND = "F4F6F8", "ECF4F8"
CAL, CALB, CALI, GEOB = "Calibri (MS)", "Calibri (MS) Bold", "Calibri (MS) Italics", "Georgia Bold"

prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]


def walk(shapes):
    for sh in shapes:
        if sh.__class__.__name__ == "GroupShape":
            yield sh
        else:
            yield sh


def kill(slide, *sids):
    want = set(sids)
    for sh in list(slide.shapes):
        if sh.shape_id in want:
            sh._element.getparent().remove(sh._element)


def flat(shape):
    """No shadow, no theme effect: the deck's shapes are flat colour only."""
    shape.shadow.inherit = False
    el = shape._element
    for child in el.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}style"):
        el.remove(child)


def rect(slide, x, y, w, h, fill):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(fill)
    r.line.fill.background()
    r.text_frame.text = ""
    flat(r)
    return r


def box(slide, x, y, w, h, runs, size=13.0, font=CAL, color=BODY, bold=None,
        italic=None, align=PP_ALIGN.LEFT, wrap=True, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para_runs in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        for text, over in para_runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = over.get("font", font)
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.italic = over.get("italic", italic)
            f.color.rgb = RGBColor.from_string(over.get("color", color))
    return tb


def band(slide, x, y, w, h, ground, accent):
    """The deck's band idiom: a flat ground with a left accent bar saying what it is."""
    rect(slide, x, y, w, h, ground)
    rect(slide, x, y, 0.14, h, accent)


# ---------------------------------------------------------------- slide 2
s2 = S(2)

# The three cards, their headers, their bullets, the stranded callout and its gloss.
kill(s2, 6, 9, 10, 11, 12, 13, 16, 17, 18, 20, 28, 29, 30, 31, 32, 33)

RX, RW = 7.30, 11.46
ATOP, BTOP, BH = 3.82, 6.52, 2.30      # two bands of equal height, right column

# ---- left: the question, stated once, large, centred against the two bands
box(s2, 1.25, 4.82, 5.4, 0.5, "THE QUESTION", size=27.99, font=GEOB, bold=True, color=TITLE)
rect(s2, 1.25, 5.39, 0.98, 0.03, RULE)
box(s2, 1.25, 5.77, 5.30, 2.2,
    [("Predict a user's ", {}),
     ("next cell_id", {"font": CALB, "bold": True, "color": TITLE}),
     (" from the beginning of their day, with a ", {}),
     ("fine-tuned language model", {"font": CALB, "bold": True, "color": TITLE}),
     (".", {})],
    size=32.0, color=BODY, spacing=1.12)

# ---- right, first band: what has to be beaten
band(s2, RX, ATOP, RW, BH, BASEGND, BASE)
box(s2, RX + 0.55, ATOP + 0.46, 5.0, 0.3, "THE COMPETITOR",
    size=17.0, font=CALB, bold=True, color=BASE, wrap=False)
box(s2, RX + 0.55, ATOP + 0.94, 3.2, 0.9, "69.2%",
    size=48.0, font=GEOB, bold=True, color=BASE, wrap=False)
box(s2, RX + 2.75, ATOP + 0.99, 8.0, 1.5,
    [[("Repeat the last cell seen.", {"font": CALB, "bold": True, "color": TITLE})],
     [("One line of code, no training. That is what the model has", {})],
     [("to beat on a raw day, and it is a hard number to beat.", {})]],
    size=19.0, color=BODY, spacing=1.14)

# ---- right, second band: how every number in the deck is reported
band(s2, RX, BTOP, RW, BH, NEUTRAL, RULE)
box(s2, RX + 0.55, BTOP + 0.52, 5.0, 0.3, "THE RULE HERE",
    size=17.0, font=CALB, bold=True, color=RULE, wrap=False)
box(s2, RX + 0.55, BTOP + 1.00, 10.5, 1.4,
    [("Every result in this deck is a ", {}),
     ("gap to that rule", {"font": CALB, "bold": True, "color": TITLE}),
     (", measured on the same users at the same scored positions. Never a bare accuracy.", {})],
    size=22.0, color=BODY, spacing=1.14)

# ---- the reading, full width, in the house band
band(s2, 1.15, 9.35, 17.61, 1.02, NEUTRAL, RULE)
box(s2, 1.70, 9.73, 16.9, 0.30,
    [("READING    ", {"font": CALB, "bold": True, "color": RULE}),
     ("The same rule scores 69.2% on a raw day and 43.1% on mobile users, so every result "
      "slide carries its corpus and its baseline, top right.", {})],
    size=15.48, color=BODY, wrap=False)

prs.save(DECK)
print("slide 2 rebuilt")
