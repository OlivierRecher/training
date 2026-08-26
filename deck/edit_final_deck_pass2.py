# -*- coding: utf-8 -*-
"""Second pass: fit and polish the edits of pass one."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK = "Présentation Stage - Finale.pptx"
TITLE, BODY, SECOND = "232959", "33363D", "47515C"
RULE, GAIN, CAVEAT, MUTED = "387190", "3F8F52", "D4760A", "8A939E"
TRACK, LINE = "E8EAEE", "C7CFD6"
CAL, CALB, CALI, GEOB = "Calibri (MS)", "Calibri (MS) Bold", "Calibri (MS) Italics", "Georgia Bold"

prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]


def walk(shapes):
    for sh in shapes:
        if sh.__class__.__name__ == "GroupShape":
            yield sh
            yield from walk(sh.shapes)
        else:
            yield sh


def sh(slide, sid, want_text=True):
    hits = [s for s in walk(slide.shapes) if s.shape_id == sid]
    if want_text:
        for s in hits:
            if s.has_text_frame and s.text_frame.paragraphs[0].runs:
                return s
    if hits:
        return hits[0]
    raise KeyError(sid)


def retext(shape, *texts, para=0):
    p = shape.text_frame.paragraphs[para]
    runs = p.runs
    while len(p.runs) < len(texts):
        p._p.append(copy.deepcopy(runs[-1]._r))
    runs = p.runs
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def move(slide, sid, **kw):
    s = sh(slide, sid)
    for k, v in kw.items():
        setattr(s, k, Inches(v))


def refill(slide, sid, rgb):
    s = sh(slide, sid)
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(rgb)


def recolor(slide, sid, rgb, run=0):
    sh(slide, sid).text_frame.paragraphs[0].runs[run].font.color.rgb = RGBColor.from_string(rgb)


def flat(shape):
    """Strip the theme style and any effect, so nothing casts a shadow."""
    shape.shadow.inherit = False
    el = shape._element
    for tag in ("{http://schemas.openxmlformats.org/presentationml/2006/main}style",):
        for child in el.findall(tag):
            el.remove(child)


def box(slide, x, y, w, h, runs, size=13.0, font=CAL, color=BODY, bold=None,
        italic=None, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, over in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = over.get("font", font)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = RGBColor.from_string(over.get("color", color))
    return tb


def chip(slide, x, y, w, h, text, fill, color, size=9.5, font=CAL):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(fill)
    r.line.fill.background()
    flat(r)
    tf = r.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    return r


# ---------------------------------------------------------------- slide 2
move(S(2), 32, top=8.50)
for s in S(2).shapes:                       # the note added in pass one
    if s.has_text_frame and s.text_frame.text.startswith("so every result slide"):
        s.top = Inches(9.26)

# ---------------------------------------------------------------- slide 6
move(S(6), 22, left=10.27, width=8.49)

# ---------------------------------------------------------------- slide 7
retext(sh(S(7), 25), "THE POINT    ",
       "Every number in this deck is a gap to the baseline, measured on the same users "
       "and the same scored positions. Where the whole sequence is scored instead of the "
       "last fifth, the slide says so, top right.")

# ---------------------------------------------------------------- slide 15
retext(sh(S(15), 4), "Group detection rate, by how much context is given")
box(S(15), 9.85, 8.05, 8.9, 0.44,
    [("FROM SLIDE 18 ON   ", {"font": CALB, "bold": True, "color": CAVEAT}),
     ("the profile no longer prepends a token: it selects which adapter answers.",
      {"font": CAL, "color": BODY})], size=13.0)

# ---------------------------------------------------------------- slide 18
move(S(18), 41, top=7.20)
for sid_text, y in ((("BEFORE",), 7.84), (("AFTER",), 8.38)):
    pass
for s in S(18).shapes:
    if s.has_text_frame and s.text_frame.text.startswith("BEFORE"):
        s.top = Inches(7.84)
    if s.has_text_frame and s.text_frame.text.startswith("<G3>"):
        s.top, s.height = Inches(8.06), Inches(0.26)
move(S(18), 47, top=8.38)
for sid in (42, 44):
    move(S(18), sid, top=8.60, height=0.24)
for sid in (48, 50, 53, 55, 58, 60, 63, 65):
    move(S(18), sid, top=8.84, height=0.28)
retext(sh(S(18), 78), "PER-GROUP WINDOWS · HOW MUCH HARDER THEY ARE")
retext(sh(S(18), 77),
       "Windows stay keyed by behavioural group, the one piece already measured as "
       "correct. Inside them persistence scores 19.0% against 46.4% outside. The global "
       "04h-06h / 18h-20h guess separated nothing at all.")
refill(S(18), 38, GAIN)        # change 3 is the structural one
refill(S(18), 74, MUTED)       # change 4 changed nothing
recolor(S(18), 82, MUTED)

# ---------------------------------------------------------------- slide 21
retext(sh(S(21), 73), "the model is the rule")
retext(sh(S(21), 78), "the entire edge, in four steps")
retext(sh(S(21), 83), "ten times more users, nothing more")
for s in S(21).shapes:
    if s.shape_type == MSO_SHAPE.RECTANGLE or s.__class__.__name__ == "Shape":
        try:
            if 4.9 < Emu(s.top).inches < 5.1:
                flat(s)
        except Exception:
            pass

# ---------------------------------------------------------------- slide 23
# Rebuild the two sequence strips so the levelled one is visibly longer.
for s in list(S(23).shapes):
    try:
        t, w = Emu(s.top).inches, Emu(s.width).inches
    except Exception:
        continue
    if abs(w - 1.22) < 0.01 and (abs(t - 6.70) < 0.02 or abs(t - 7.98) < 0.02):
        s._element.getparent().remove(s._element)

NAT = ["BKVCHE1", "BKVCHE1", "UKVSUM101", "BKVPAN1", "BKVPAN1", "DSOTKB2"]
for i, c in enumerate(NAT):
    chip(S(23), 1.25 + i * 1.22, 6.70, 1.15, 0.30, c, TRACK, BODY, size=9.0)
x = 1.25
for c in NAT:
    chip(S(23), x, 7.98, 0.62, 0.30, c[3:], "FCEEDC", BODY, size=7.0)
    x += 0.66
    chip(S(23), x, 7.98, 0.62, 0.30, c[3:], TRACK, MUTED, size=7.0)
    x += 0.66
box(S(23), x + 0.06, 8.02, 0.9, 0.24, "…", size=14.0, font=CALB, bold=True, color=MUTED, wrap=False)
for s in S(23).shapes:
    if s.has_text_frame and s.text_frame.text.startswith("↓"):
        retext(s, "↓   ", "each record repeated a few times, spread evenly along the day, "
                          "never piled up at one end")
    if s.has_text_frame and s.text_frame.text.startswith("exactly 200"):
        s.top = Inches(8.34)

prs.save(DECK)
print("saved")
