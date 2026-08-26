# -*- coding: utf-8 -*-
"""Third pass: two misplacements and the sequence-strip labels."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK = "Présentation Stage - Finale.pptx"
BODY, SECOND, MUTED, CAVEAT = "33363D", "47515C", "8A939E", "D4760A"
TRACK = "E8EAEE"
CAL, CALB, CALI = "Calibri (MS)", "Calibri (MS) Bold", "Calibri (MS) Italics"

prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]


def walk(shapes):
    for sh in shapes:
        if sh.__class__.__name__ == "GroupShape":
            yield sh
            yield from walk(sh.shapes)
        else:
            yield sh


def sh(slide, sid):
    hits = [s for s in walk(slide.shapes) if s.shape_id == sid]
    for s in hits:
        if s.has_text_frame and s.text_frame.paragraphs[0].runs:
            return s
    return hits[0]


def flat(shape):
    shape.shadow.inherit = False
    el = shape._element
    for child in el.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}style"):
        el.remove(child)


def box(slide, x, y, w, h, runs, size=13.0, font=CAL, color=BODY, bold=None,
        italic=None, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, over in (runs if not isinstance(runs, str) else [(runs, {})]):
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = over.get("font", font)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = RGBColor.from_string(over.get("color", color))
    return tb


def chip(slide, x, y, w, h, text, fill, color, size=9.0):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(fill)
    r.line.fill.background()
    flat(r)
    tf = r.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = CAL
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    return r


# slide 15: air under the arrow line
for s in S(15).shapes:
    if s.has_text_frame and s.text_frame.text.startswith("FROM SLIDE 18 ON"):
        s.top = Inches(8.20)

# slide 18: put box 2's BEFORE label back where it belongs, and match the AFTER label
sh(S(18), 28).top = Inches(4.90)
old = sh(S(18), 47)
old._element.getparent().remove(old._element)
box(S(18), 1.55, 8.38, 7.68, 0.22,
    [("AFTER", {"font": CALB, "bold": True, "color": "3F8F52"}),
     ("   the profile chooses which weights answer", {"font": CAL, "color": SECOND})],
    size=9.99, wrap=False)

# slide 23: full cell_id on the levelled strip, and a legend for the shading
for s in list(S(23).shapes):
    try:
        t, w = Emu(s.top).inches, Emu(s.width).inches
    except Exception:
        continue
    if abs(w - 0.62) < 0.01 and abs(t - 7.98) < 0.02:
        s._element.getparent().remove(s._element)
NAT = ["BKVCHE1", "BKVCHE1", "UKVSUM101", "BKVPAN1", "BKVPAN1", "DSOTKB2"]
x = 1.25
for c in NAT:
    chip(S(23), x, 7.98, 0.62, 0.30, c, "FCEEDC", BODY, size=7.0)
    chip(S(23), x + 0.66, 7.98, 0.62, 0.30, c, TRACK, MUTED, size=7.0)
    x += 1.32
for s in S(23).shapes:
    if s.has_text_frame and s.text_frame.text.startswith("exactly 200"):
        p = s.text_frame.paragraphs[0]
        p.runs[0].text = ("exactly 200 records for every user, so every user weighs 0.25% of "
                          "the training. Solid is the original record, pale is its copy.")

prs.save(DECK)
print("saved")
